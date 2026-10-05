import csv
import io
from typing import AsyncGenerator
from openpyxl import load_workbook
from email_validator import validate_email, EmailNotValidError


def _decode_csv_content(file_content: bytes) -> str:
    """Decode CSV content with fallback encoding detection."""
    for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        try:
            return file_content.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_content.decode("utf-8-sig", errors="replace")


def _detect_csv_delimiter(text: str) -> str:
    """Detect delimiter from sample text."""
    try:
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        return dialect.delimiter
    except Exception:
        return ","


def parse_csv_headers(file_content: bytes) -> list[str]:
    """Extract column headers from CSV content with stripped whitespace and quotes."""
    text = _decode_csv_content(file_content)
    delimiter = _detect_csv_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    headers = next(reader, None)
    if not headers:
        return []
    return [str(h).strip().strip("'\"") for h in headers if h is not None]


def parse_excel_headers(file_content: bytes) -> list[str]:
    """Extract column headers from Excel content."""
    wb = load_workbook(filename=io.BytesIO(file_content), read_only=True)
    ws = wb.active
    headers = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)):
        headers.append(str(cell).strip() if cell else "")
    wb.close()
    return headers


def parse_csv_rows(file_content: bytes, batch_size: int = 500):
    """Parse CSV rows in batches. Returns generator of row dicts with stripped keys and values."""
    text = _decode_csv_content(file_content)
    delimiter = _detect_csv_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_headers = next(reader, None)
    if not raw_headers:
        return

    # Normalize headers exactly as parse_csv_headers does
    headers = [str(h).strip().strip("'\"") for h in raw_headers if h is not None]

    batch = []
    for values in reader:
        if not values or not any(v.strip() for v in values if isinstance(v, str)):
            continue
        row_dict = {}
        for i, val in enumerate(values):
            if i < len(headers):
                row_dict[headers[i]] = str(val).strip() if val is not None else ""
        batch.append(row_dict)
        if len(batch) >= batch_size:
            yield batch
            batch = []

    if batch:
        yield batch


def parse_excel_rows(file_content: bytes, batch_size: int = 500):
    """Parse Excel rows in batches. Returns generator of row dicts."""
    wb = load_workbook(filename=io.BytesIO(file_content), read_only=True)
    ws = wb.active

    # Get headers from first row
    headers = []
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = [str(cell).strip() if cell else f"col_{i}" for i, cell in enumerate(row)]
        break

    batch = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = {}
        for i, value in enumerate(row):
            if i < len(headers):
                row_dict[headers[i]] = str(value).strip() if value is not None else ""
        batch.append(row_dict)
        if len(batch) >= batch_size:
            yield batch
            batch = []

    if batch:
        yield batch
    wb.close()


def validate_email_address(email: str) -> bool:
    """Validate an email address format."""
    try:
        validate_email(email, check_deliverability=False)
        return True
    except EmailNotValidError:
        return False


def count_rows(file_content: bytes, file_type: str) -> int:
    """Count total data rows in the file."""
    if file_type == "csv":
        text = _decode_csv_content(file_content)
        delimiter = _detect_csv_delimiter(text)
        return max(sum(1 for _ in csv.reader(io.StringIO(text), delimiter=delimiter)) - 1, 0)
    else:
        wb = load_workbook(filename=io.BytesIO(file_content), read_only=True)
        ws = wb.active
        count = ws.max_row - 1 if ws.max_row else 0
        wb.close()
        return max(count, 0)
